import streamlit as st
import streamlit.components.v1 as components
import fitz  # PyMuPDF
import re
import io
import json
import zipfile
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Universal Operations & e-Invoice Suite", page_icon="⚡", layout="wide")

st.markdown("""
    <div style="text-align: center; padding: 15px 0 20px 0;">
        <h2 style="color: #FFFFFF; margin-bottom: 6px;">⚡ Universal All-Invoice Operations Suite</h2>
        <p style="color: #94a3b8; font-size: 14px;">Original Product Description Restored | Col AR Blank | Exact Discount (Col BO) & Round Off</p>
    </div>
""", unsafe_allow_html=True)

uploaded_invoices = st.file_uploader("Upload Invoices (PDF) - Single ya Bulk", type=["pdf"], accept_multiple_files=True)

STATE_CODE_MAP = {
    "01": "JAMMU AND KASHMIR", "02": "HIMACHAL PRADESH", "03": "PUNJAB", "04": "CHANDIGARH",
    "05": "UTTARAKHAND", "06": "HARYANA", "07": "DELHI", "08": "RAJASTHAN",
    "09": "UTTAR PRADESH", "10": "BIHAR", "11": "SIKKIM", "12": "ARUNACHAL PRADESH",
    "13": "NAGALAND", "14": "MANIPUR", "15": "MIZORAM", "16": "TRIPURA",
    "17": "MEGHALAYA", "18": "ASSAM", "19": "WEST BENGAL", "20": "JHARKHAND",
    "21": "ODISHA", "22": "CHATTISGARH", "23": "MADHYA PRADESH", "24": "GUJARAT",
    "26": "DADRA AND NAGAR HAVELI AND DAMAN AND DIU", "27": "MAHARASHTRA", "29": "KARNATAKA",
    "30": "GOA", "31": "LAKSHADWEEP", "32": "KERALA", "33": "TAMIL NADU",
    "34": "PUDUCHERRY", "35": "ANDAMAN AND NICOBAR ISLANDS", "36": "TELANGANA", "37": "ANDHRA PRADESH",
    "38": "LADAKH"
}

STATE_NAME_TO_CODE = {v.upper(): k for k, v in STATE_CODE_MAP.items()}

MAJOR_CITIES = [
    "New Delhi", "Delhi", "Varanasi", "Lucknow", "Jaipur", "Gurgaon", "Gurugram", 
    "Noida", "Ghaziabad", "Kanpur", "Bengaluru", "Bangalore", "Mumbai", "Pune", 
    "Kolkata", "Ahmedabad", "Patna", "Ranchi", "Chandigarh", "Faridabad", "Agra", "Meerut"
]

def clean_extracted_city(text, state_name):
    m = re.search(r"([A-Za-z\s]+),\s*(?:[A-Za-z\s]+)[\-\s]*[0-9]{6}", text)
    if m:
        c_cand = m.group(1).strip()
        parts = [p.strip() for p in c_cand.split(",") if p.strip()]
        last_p = parts[-1]
        if len(last_p) > 2 and not re.search(r"(road|street|nagar|colony|floor|block|house|marg)", last_p, re.I):
            return last_p.title()

    for city in MAJOR_CITIES:
        if re.search(rf"\b{city}\b", text, re.I):
            return city

    return state_name.title() if state_name else "Delhi"

def fmt_dec(val):
    try:
        f = float(str(val).replace(",", "").strip())
        return f"{f:.2f}"
    except (ValueError, TypeError):
        return "0.00"

def clean_description_completely(desc, hsn_code=None):
    if not desc:
        return ""
    if hsn_code:
        desc = re.sub(rf"\b{re.escape(str(hsn_code))}\b", "", desc)
    desc = re.sub(r"\b\d{6,8}\b", "", desc)
    desc = re.sub(r"\s+", " ", desc).strip()
    return desc

# ----------------- ORIGINAL STABLE METADATA EXTRACTION -----------------

def extract_metadata_from_doc(doc):
    full_text = ""
    for p in doc:
        full_text += p.get_text() + "\n"

    page0 = doc[0]
    words = page0.get_text("words")

    ext_order_id = None
    ext_match = re.search(r"Extern(?:al)?\s*Order\s*(?:No\.?|ID)?\s*[:\s]*([0-9\sA-Za-z\-]+?)(?=\n|Invoice|Date|$)", full_text, re.IGNORECASE)
    if ext_match:
        ext_order_id = "".join(ext_match.group(1).split())
    if not ext_order_id:
        digits_match = re.findall(r"\b(49\d{8,14}|5\d{8,14}|\d{12,18})\b", full_text)
        if digits_match:
            ext_order_id = digits_match[0]
    if not ext_order_id:
        ext_order_id = "ExtOrder"

    inv_match = re.search(r"Invoice\s*No\s*[:\s]*([A-Za-z0-9\-_/]+)", full_text, re.IGNORECASE)
    invoice_no = inv_match.group(1).strip() if inv_match else "INV-001"

    date_match = re.search(r"Invoice\s*Date\s*[:\s]*([A-Za-z0-9,\s\.\-\/]+?)(?=\n|Ship\s*Date|Order|$)", full_text, re.IGNORECASE)
    raw_date = date_match.group(1).strip() if date_match else "Date"

    clean_date = raw_date
    std_date_for_json = datetime.now().strftime("%d/%m/%Y")
    for fmt in ("%b %d, %Y", "%B %d, %Y", "%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            d_obj = datetime.strptime(raw_date, fmt)
            clean_date = d_obj.strftime("%d-%m-%Y")
            std_date_for_json = d_obj.strftime("%d/%m/%Y")
            break
        except ValueError:
            pass

    clean_ext = re.sub(r'[^A-Za-z0-9\-_]', '', ext_order_id)
    clean_inv = re.sub(r'[^A-Za-z0-9\-_]', '', invoice_no)
    clean_dt = re.sub(r'[^A-Za-z0-9\-_]', '', clean_date)
    pdf_filename = f"{clean_ext}_{clean_inv}_{clean_dt}.pdf"

    sold_header = [w for w in words if "sold" in w[4].lower()]
    bill_header = [w for w in words if "billing" in w[4].lower()]
    table_header = [w for w in words if w[4].lower() in ["s.no", "sl.no", "item", "hsn"]]

    y_sold = sold_header[0][1] if sold_header else 80
    y_bill = bill_header[0][1] if bill_header else 215
    y_table = table_header[0][1] if table_header else 340
    split_x = 300

    seller_rect = fitz.Rect(20, y_sold, split_x, y_bill)
    seller_text = page0.get_text("text", clip=seller_rect).strip()
    s_lines = [l.strip() for l in seller_text.split("\n") if l.strip() and not re.search(r"sold\s*by", l, re.I)]
    seller_name = s_lines[0] if s_lines else "ROMSONS PRIME PRIVATE LIMITED"

    s_gst = re.search(r"GSTIN\s*[:\s]*([0-9]{2}[A-Z]{5}[0-9]{4}[A-Z]{1}[1-9A-Z]{1}Z[0-9A-Z]{1})", seller_text, re.I)
    seller_gstin = s_gst.group(1).strip() if s_gst else "07AALCR5906L1ZW"

    s_pin = re.findall(r"\b[1-9][0-9]{5}\b", seller_text)
    seller_pin = int(s_pin[0]) if s_pin else 110028
    seller_state_code = seller_gstin[:2]
    seller_state_name = STATE_CODE_MAP.get(seller_state_code, "DELHI")
    seller_loc = clean_extracted_city(seller_text, seller_state_name)

    seller_addr_candidates = [l for l in s_lines[1:] if not re.search(r"(@|contact|phone|mob|gstin|pan|india)", l, re.I)]
    seller_addr = ", ".join(seller_addr_candidates[:2]) if seller_addr_candidates else "Ground Floor, B-248 B Block, Naraina Industrial Area, Phase 1"

    billing_rect = fitz.Rect(20, y_bill, split_x, y_table)
    billing_text = page0.get_text("text", clip=billing_rect).strip()
    b_lines = [l.strip() for l in billing_text.split("\n") if l.strip() and not re.search(r"billing\s*addr", l, re.I)]

    buyer_name = b_lines[0] if b_lines else "BUYER ENTERPRISE"
    b_addr_candidates = [l for l in b_lines[1:] if not re.search(r"(@|contact|phone|mob|gstin|pan|india)", l, re.I)]
    buyer_addr1 = ", ".join(b_addr_candidates[:2]) if b_addr_candidates else (b_lines[1] if len(b_lines) > 1 else "Commercial Facility")

    em_m = re.search(r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+", billing_text)
    buyer_email = em_m.group(0).strip() if em_m else ""

    ph_m = re.search(r"(?:Contact|Phone|Mob)?\s*[:\s]*([6-9]\d{9})", billing_text, re.IGNORECASE)
    buyer_phone = ph_m.group(1).strip() if ph_m else ""

    b_gst_m = re.search(r"GSTIN\s*[:\s]*([0-9]{2}[A-Z]{5}[0-9]{4}[A-Z]{1}[1-9A-Z]{1}Z[0-9A-Z]{1})", billing_text, re.IGNORECASE)
    buyer_gstin = b_gst_m.group(1).strip() if b_gst_m else ""

    if not buyer_gstin:
        all_gstins = re.findall(r"\b\d{2}[A-Z]{5}\d{4}[A-Z]{1}[1-9A-Z]{1}Z[0-9A-Z]{1}\b", full_text)
        buyer_gstin = all_gstins[1] if len(all_gstins) > 1 else (all_gstins[0] if all_gstins else "07AACCT5680N1ZS")

    buyer_state_code = buyer_gstin[:2]
    buyer_state_name = STATE_CODE_MAP.get(buyer_state_code, "DELHI")

    b_pin_m = re.findall(r"\b[1-9][0-9]{5}\b", billing_text)
    buyer_pin = int(b_pin_m[0]) if b_pin_m else int(buyer_state_code + "0001")
    buyer_loc = clean_extracted_city(billing_text, buyer_state_name)

    shipping_rect = fitz.Rect(split_x + 5, y_bill, 585, y_table)
    shipping_text = page0.get_text("text", clip=shipping_rect).strip()
    shp_lines = [l.strip() for l in shipping_text.split("\n") if l.strip() and not re.search(r"shipping\s*addr", l, re.I)]

    ship_name = shp_lines[0] if shp_lines else buyer_name
    s_addr_candidates = [l for l in shp_lines[1:] if not re.search(r"(@|contact|phone|mob|gstin|pan|india)", l, re.I)]
    ship_addr1 = ", ".join(s_addr_candidates[:2]) if s_addr_candidates else (shp_lines[1] if len(shp_lines) > 1 else buyer_addr1)

    shp_gst_m = re.search(r"GSTIN\s*[:\s]*([0-9]{2}[A-Z]{5}[0-9]{4}[A-Z]{1}[1-9A-Z]{1}Z[0-9A-Z]{1})", shipping_text, re.IGNORECASE)
    ship_gstin = shp_gst_m.group(1).strip() if shp_gst_m else buyer_gstin
    ship_state_code = ship_gstin[:2]
    ship_state_name = STATE_CODE_MAP.get(ship_state_code, buyer_state_name)

    shp_pin_m = re.findall(r"\b[1-9][0-9]{5}\b", shipping_text)
    ship_pin = int(shp_pin_m[0]) if shp_pin_m else buyer_pin
    ship_loc = clean_extracted_city(shipping_text, ship_state_name)

    clean_b_str = re.sub(r'[^A-Za-z0-9]', '', f"{buyer_addr1}{buyer_pin}").lower()
    clean_s_str = re.sub(r'[^A-Za-z0-9]', '', f"{ship_addr1}{ship_pin}").lower()
    shipping_is_different = (clean_b_str != clean_s_str) and (buyer_pin != ship_pin or buyer_addr1 != ship_addr1)

    # Place of Supply Detection
    pos_match = re.search(r"Place\s*of\s*Supply\s*[:\s]*([A-Za-z0-9\s\-]+?)(?=\n|GST|State|$)", full_text, re.IGNORECASE)
    pos_raw = pos_match.group(1).strip() if pos_match else ""
    pos_digits = re.findall(r"\b\d{2}\b", pos_raw)

    if pos_digits:
        buyer_pos_code = pos_digits[0]
        pos_state_name = STATE_CODE_MAP.get(buyer_pos_code, buyer_state_name)
    else:
        matched_code = None
        for s_name, s_code in STATE_NAME_TO_CODE.items():
            if s_name in pos_raw.upper():
                matched_code = s_code
                pos_state_name = s_name.title()
                break
        buyer_pos_code = matched_code if matched_code else buyer_state_code
        pos_state_name = STATE_CODE_MAP.get(buyer_pos_code, buyer_state_name)

    is_blinkit = bool(re.search(r"blink\s*commerce|blinkit", f"{buyer_name} {full_text}", re.IGNORECASE))

    # Accurate Tax Type Detection directly from invoice
    has_cgst_header = bool(re.search(r"\bCGST\b", full_text, re.I))
    cgst_vals = re.findall(r"CGST[\s\S]*?(\d+\.\d{2})", full_text)
    has_positive_cgst = any(float(v) > 0 for v in cgst_vals[:4]) if cgst_vals else False

    if has_cgst_header and has_positive_cgst:
        is_interstate = False
    else:
        is_interstate = (seller_state_code != buyer_pos_code)

    inv_total_m = re.search(r"(?:Total\s*Invoice\s*Value|Invoice\s*Total|Grand\s*Total|Total\s*Amount)[:\s]*[₹\s]*([0-9,]+\.\d{2})", full_text, re.I)
    printed_grand_total = float(inv_total_m.group(1).replace(",", "")) if inv_total_m else None

    return {
        "invoice_no": invoice_no,
        "doc_date": std_date_for_json,
        "clean_date": clean_date,
        "pdf_filename": pdf_filename,
        "seller_name": seller_name,
        "seller_gstin": seller_gstin,
        "seller_pin": seller_pin,
        "seller_addr": seller_addr,
        "seller_loc": seller_loc,
        "seller_state_code": seller_state_code,
        "seller_state_name": seller_state_name,
        "buyer_name": buyer_name,
        "buyer_trade_name": buyer_name,
        "buyer_addr1": buyer_addr1,
        "buyer_loc": buyer_loc,
        "buyer_pin": buyer_pin,
        "buyer_phone": buyer_phone,
        "buyer_email": buyer_email,
        "buyer_gstin": buyer_gstin,
        "buyer_state_code": buyer_state_code,
        "buyer_state_name": buyer_state_name,
        "buyer_pos": buyer_pos_code,
        "pos_state_name": pos_state_name,
        "is_interstate": is_interstate,
        "shipping_is_different": shipping_is_different,
        "ship_name": ship_name,
        "ship_addr1": ship_addr1,
        "ship_loc": ship_loc,
        "ship_pin": ship_pin,
        "ship_gstin": ship_gstin,
        "ship_state_name": ship_state_name,
        "ship_state_code": ship_state_code,
        "is_blinkit": is_blinkit,
        "printed_grand_total": printed_grand_total
    }

# ----------------- ORIGINAL ACCURATE TABLE PARSER RESTORED -----------------

def overwrite_area(page, rect, new_text, font_size=7):
    pad_rect = fitz.Rect(rect.x0 - 2, rect.y0 - 1, rect.x1 + 2, rect.y1 + 1)
    page.draw_rect(pad_rect, color=None, fill=(1, 1, 1))
    page.insert_text((rect.x0, rect.y1 - 1.2), str(new_text), fontsize=font_size, fontname="helv", color=(0, 0, 0))

def process_and_reconcile_pdf(pdf_bytes):
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    meta = extract_metadata_from_doc(doc)
    is_blinkit = meta["is_blinkit"]
    line_items = []

    for page in doc:
        words = page.get_text("words")

        sno_anchor_x = 60
        hsn_min_x = 160
        hsn_max_x = 240
        qty_left = 240
        price_left = 300
        disc_left = 360
        tax_left = 410

        for w in words:
            wt = w[4].lower().strip()
            if wt in ["s.no", "sl.no", "s.no."]:
                sno_anchor_x = w[2] + 6
            elif wt == "hsn":
                hsn_min_x = w[0] - 12
                hsn_max_x = w[2] + 30
            elif wt in ["qty", "quantity"]:
                qty_left = w[0] - 8
            elif "price" in wt or "unit" in wt or "rate" in wt:
                price_left = w[0] - 8
            elif "disc" in wt:
                disc_left = w[0] - 8
            elif "taxable" in wt:
                tax_left = w[0] - 8

        sno_candidates = []
        for w in words:
            val = w[4].strip()
            if val.isdigit() and 1 <= int(val) <= 99:
                if w[0] <= sno_anchor_x and w[1] > 200:
                    if not any(abs(c["center_y"] - ((w[1]+w[3])/2)) < 8 for c in sno_candidates):
                        sno_candidates.append({
                            "sno": int(val),
                            "y0": w[1],
                            "y1": w[3],
                            "center_y": (w[1] + w[3]) / 2
                        })

        sno_candidates.sort(key=lambda x: x["center_y"])

        table_bottom_y = page.rect.height - 40
        total_indicators = [w for w in words if w[1] > 300 and any(k in w[4].lower() for k in ["total", "subtotal", "taxable value", "gst summary", "amount in words"])]
        if total_indicators:
            table_bottom_y = min([w[1] for w in total_indicators]) - 4

        if sno_candidates:
            for idx, item_anchor in enumerate(sno_candidates):
                y_top = item_anchor["y0"] - 3
                if idx + 1 < len(sno_candidates):
                    y_bottom = sno_candidates[idx + 1]["y0"] - 2
                else:
                    y_bottom = min(item_anchor["y1"] + 110, table_bottom_y)

                row_words = [w for w in words if y_top <= ((w[1] + w[3]) / 2) <= y_bottom]

                # 1. HSN Code
                valid_hsns = [
                    w[4].replace(",", "").strip() for w in row_words 
                    if (hsn_min_x - 20) <= w[0] <= (hsn_max_x + 20) and w[4].replace(",", "").strip().isdigit() and len(w[4].replace(",", "").strip()) in [6, 7, 8]
                ]
                hsn_code = valid_hsns[0] if valid_hsns else "90183990"

                actual_hsn_x = hsn_min_x
                for w in row_words:
                    if w[4].replace(",", "").strip() == hsn_code:
                        actual_hsn_x = min(actual_hsn_x, w[0])
                        break

                # 2. Original Exact Product Description
                item_right_boundary = actual_hsn_x - 3
                item_box_rect = fitz.Rect(sno_anchor_x - 10, y_top, item_right_boundary, y_bottom)
                raw_item_text = page.get_text("text", clip=item_box_rect).strip()
                lines = [l.strip() for l in raw_item_text.split("\n") if l.strip()]
                raw_desc = " ".join(lines).strip()

                pure_desc = clean_description_completely(raw_desc, hsn_code=hsn_code)
                if not pure_desc:
                    pure_desc = f"Item {item_anchor['sno']}"

                if any(k in pure_desc.lower() for k in ["cod charge", "shipping charge"]):
                    continue

                factor = 1
                if is_blinkit:
                    upper_desc = pure_desc.upper()
                    if "DISPO" in upper_desc:
                        factor = 50
                    elif "COMFIT" in upper_desc:
                        factor = 25

                final_q = None
                final_p = None
                printed_taxable = None
                printed_discount = 0.0

                # 3. Quantity
                qty_box_rect = fitz.Rect(qty_left - 10, y_top, price_left - 2, y_bottom)
                for w in words:
                    if qty_box_rect.contains(fitz.Point((w[0]+w[2])/2, (w[1]+w[3])/2)):
                        val = w[4].replace(",", "").strip()
                        rect = fitz.Rect(w[0], w[1], w[2], w[3])
                        if val.isdigit() and len(val) <= 6:
                            orig_q = int(val)
                            if factor > 1 and orig_q >= factor:
                                final_q = orig_q // factor
                                overwrite_area(page, rect, f"{final_q}")
                            else:
                                final_q = orig_q
                            break

                # 4. Unit Price
                price_box_rect = fitz.Rect(price_left - 10, y_top, disc_left - 2, y_bottom)
                for w in words:
                    if price_box_rect.contains(fitz.Point((w[0]+w[2])/2, (w[1]+w[3])/2)):
                        val = w[4].replace(",", "").strip()
                        rect = fitz.Rect(w[0], w[1], w[2], w[3])
                        if re.match(r"^\d+(\.\d{1,2})?$", val):
                            orig_p = float(val)
                            if orig_p > 0.00:
                                if factor > 1:
                                    final_p = round(orig_p * factor, 2)
                                    overwrite_area(page, rect, f"{final_p:.2f}")
                                else:
                                    final_p = orig_p
                                break

                # 5. Discount Column
                disc_box_rect = fitz.Rect(disc_left - 10, y_top, tax_left - 2, y_bottom)
                raw_disc_text = page.get_text("text", clip=disc_box_rect)
                disc_m = re.findall(r"\b\d+(?:\.\d{1,2})?\b", raw_disc_text)
                if disc_m:
                    printed_discount = float(disc_m[0])

                # 6. Taxable Value Column
                tax_box_rect = fitz.Rect(tax_left - 10, y_top, tax_left + 90, y_bottom)
                raw_tax_text = page.get_text("text", clip=tax_box_rect)
                tax_m = re.findall(r"\b\d+[.,]\d{2}\b", raw_tax_text)
                if tax_m:
                    printed_taxable = float(tax_m[0].replace(",", ""))

                if final_q is not None and final_p is not None and final_q > 0:
                    gross_amt = round(final_q * final_p, 2)

                    if printed_taxable is not None and printed_taxable > 0:
                        taxable_val = printed_taxable
                        if gross_amt > taxable_val and printed_discount == 0.0:
                            printed_discount = round(gross_amt - taxable_val, 2)
                    else:
                        taxable_val = round(gross_amt - printed_discount, 2)

                    unit_label = "BOX" if (is_blinkit and factor > 1) else "PAC"
                    line_items.append({
                        "sno": len(line_items) + 1,
                        "desc": pure_desc,
                        "hsn": hsn_code,
                        "qty": final_q,
                        "unit": unit_label,
                        "unit_price": final_p,
                        "gross_amt": gross_amt,
                        "discount": printed_discount,
                        "taxable_val": taxable_val,
                        "gst_rate": 5.0
                    })

        if is_blinkit:
            for target in ["UOM-PC", "UOM-IBOX", "UOM-PCS", "UOM-BOX"]:
                for inst in page.search_for(target):
                    overwrite_area(page, inst, "UOM-BOX")

            for s_inst in page.search_for("S"):
                if 140 <= s_inst.x0 <= 260:
                    page.draw_rect(s_inst, color=None, fill=(1, 1, 1))

    out_pdf_buf = io.BytesIO()
    doc.save(out_pdf_buf)
    doc.close()
    out_pdf_buf.seek(0)

    meta["line_items"] = line_items
    return out_pdf_buf, meta

# ----------------- JSON v1.01 BUILDER -----------------

def build_einvoice_json_v101(data):
    is_interstate = data["is_interstate"]
    seller_state_code = data["seller_state_code"]
    buyer_pos = data["buyer_pos"]

    item_list = []
    tot_taxable = 0.0
    tot_cgst = 0.0
    tot_sgst = 0.0
    tot_igst = 0.0

    for idx, item in enumerate(data["line_items"], 1):
        qty = item["qty"]
        unit_price = item["unit_price"]
        unit_label = item.get("unit", "BOX")
        taxable_amt = item.get("taxable_val", round(qty * unit_price, 2))
        tot_taxable += taxable_amt

        gst_rate = item.get("gst_rate", 5.0)

        if is_interstate:
            igst_amt = round((taxable_amt * gst_rate) / 100, 2)
            cgst_amt = 0.0
            sgst_amt = 0.0
            tot_igst += igst_amt
        else:
            cgst_amt = round((taxable_amt * (gst_rate / 2)) / 100, 2)
            sgst_amt = round((taxable_amt * (gst_rate / 2)) / 100, 2)
            igst_amt = 0.0
            tot_cgst += cgst_amt
            tot_sgst += sgst_amt

        tot_item_val = round(taxable_amt + cgst_amt + sgst_amt + igst_amt, 2)

        item_list.append({
            "SlNo": str(idx),
            "PrdDesc": item["desc"][:250],
            "IsServc": "N",
            "HsnCd": item["hsn"],
            "Qty": qty,
            "Unit": unit_label,
            "UnitPrice": round(unit_price, 2),
            "TotAmt": round(item.get("gross_amt", taxable_amt), 2),
            "Discount": round(item.get("discount", 0.0), 2),
            "AssAmt": round(taxable_amt, 2),
            "GstRt": gst_rate,
            "IgstAmt": igst_amt,
            "CgstAmt": cgst_amt,
            "SgstAmt": sgst_amt,
            "CesRt": 0.0,
            "CesAmt": 0.0,
            "TotItemVal": tot_item_val
        })

    exact_calc_total = round(tot_taxable + tot_cgst + tot_sgst + tot_igst, 2)
    final_rounded_total = round(exact_calc_total)
    round_off_amt = round(final_rounded_total - exact_calc_total, 2)

    ship_dtls = None
    if data.get("shipping_is_different", False):
        ship_dtls = {
            "Gstin": data["ship_gstin"],
            "LglNm": data["ship_name"],
            "TrdNm": data["ship_name"],
            "Addr1": data["ship_addr1"][:100],
            "Loc": data["ship_loc"],
            "Pin": data["ship_pin"],
            "Stcd": data["ship_state_code"]
        }

    return {
        "Version": "1.01",
        "TranDtls": {
            "TaxSch": "GST",
            "SupTyp": "B2B",
            "RegRev": "N",
            "EcmGstin": None,
            "IgstOnIntra": "N"
        },
        "DocDtls": {
            "Typ": "INV",
            "No": data["invoice_no"],
            "Dt": data["doc_date"]
        },
        "SellerDtls": {
            "Gstin": data["seller_gstin"],
            "LglNm": data["seller_name"],
            "TrdNm": data["seller_name"],
            "Addr1": f"Warehouse Facility, {data['seller_loc']}",
            "Loc": data["seller_loc"],
            "Pin": data["seller_pin"],
            "Stcd": seller_state_code
        },
        "BuyerDtls": {
            "Gstin": data["buyer_gstin"],
            "LglNm": data["buyer_name"],
            "TrdNm": data["buyer_trade_name"],
            "Pos": buyer_pos,
            "Addr1": data["buyer_addr1"][:100],
            "Loc": data["buyer_loc"],
            "Pin": data["buyer_pin"],
            "Stcd": buyer_pos
        },
        "DispDtls": None,
        "ShipDtls": ship_dtls,
        "ItemList": item_list,
        "ValDtls": {
            "AssVal": round(tot_taxable, 2),
            "CgstVal": round(tot_cgst, 2),
            "SgstVal": round(tot_sgst, 2),
            "IgstVal": round(tot_igst, 2),
            "CesVal": 0.0,
            "StCesVal": 0.0,
            "RndOffAmt": round_off_amt,
            "TotInvVal": round(final_rounded_total, 2)
        }
    }

# ----------------- 100% ACCURATE EXCEL GENERATOR (RAW BYTES) -----------------

def generate_official_nic_v101_excel_bytes(data_list):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "eInvoice"
    ws.views.sheetView[0].showGridLines = True

    c_supply = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    c_doc = PatternFill(start_color="FDE9D9", end_color="FDE9D9", fill_type="solid")
    c_buyer = PatternFill(start_color="B8CCE4", end_color="B8CCE4", fill_type="solid")
    c_disp = PatternFill(start_color="FDE9D9", end_color="FDE9D9", fill_type="solid")
    c_ship = PatternFill(start_color="B8CCE4", end_color="B8CCE4", fill_type="solid")
    c_item = PatternFill(start_color="D8E4BC", end_color="D8E4BC", fill_type="solid")
    c_val = PatternFill(start_color="B8CCE4", end_color="B8CCE4", fill_type="solid")
    c_export = PatternFill(start_color="F2DCDB", end_color="F2DCDB", fill_type="solid")
    c_eway = PatternFill(start_color="E6B8B7", end_color="E6B8B7", fill_type="solid")

    font_title = Font(name="Arial", size=18, bold=True)
    font_btn = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    font_sec = Font(name="Arial", size=10, bold=True)
    font_col = Font(name="Arial", size=9, bold=True)
    font_data = Font(name="Arial", size=9)

    thin_border = Border(
        left=Side(style='thin', color='B0B0B0'),
        right=Side(style='thin', color='B0B0B0'),
        top=Side(style='thin', color='B0B0B0'),
        bottom=Side(style='thin', color='B0B0B0')
    )

    ws.merge_cells("H1:L2")
    ws["H1"] = "E-Invoice System"
    ws["H1"].font = font_title
    ws["H1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("M1:N2")
    ws["M1"] = "Validate"
    ws["M1"].fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    ws["M1"].font = font_btn
    ws["M1"].alignment = Alignment(horizontal="center", vertical="center")

    sections = [
        ("Supply Details", 1, 4, c_supply),
        ("Document Details", 5, 7, c_doc),
        ("Buyer Details", 8, 18, c_buyer),
        ("Dispatch Details", 19, 24, c_disp),
        ("Shipping Details", 25, 32, c_ship),
        ("Product Details", 33, 57, c_item),
        ("Value Details", 61, 70, c_val),
        ("Export Details", 71, 77, c_export),
        ("E-way-bill Details", 78, 85, c_eway)
    ]

    for name, sc, ec, fill in sections:
        ws.merge_cells(start_row=3, start_column=sc, end_row=3, end_column=ec)
        c = ws.cell(row=3, column=sc, value=name)
        c.fill = fill
        c.font = font_sec
        c.alignment = Alignment(horizontal="center", vertical="center")
        for i in range(sc, ec + 1):
            ws.cell(row=3, column=i).border = thin_border

    col_names = [
        # Supply Details (1 to 4: A-D)
        "Supply Type Code *", "Reverse Charge", "e-Comm GSTIN", "Igst On Intra",
        # Document Details (5 to 7: E-G)
        "Document Type *", "Document Number *", "Document Date (DD/MM/YYYY) *",
        # Buyer Details (8 to 18: H-R)
        "Buyer GSTIN *", "Buyer Legal Name *", "Buyer Trade Name", "Buyer POS *",
        "Buyer Addr1 *", "Buyer Addr2", "Buyer Location *", "Buyer Pin Code *",
        "Buyer State *", "Buyer Phone Number", "Buyer Email Id",
        # Dispatch Details (19 to 24: S-X)
        "Dispatch Name", "Dispatch Addr1", "Dispatch Addr2", "Dispatch Location", "Dispatch Pin Code", "Dispatch State",
        # Shipping Details (25 to 32: Y-AF)
        "Shipping GSTIN", "Shipping Legal Name", "Shipping Trade Name", "Shipping Addr1", "Shipping Addr2", "Shipping Location", "Shipping Pin Code", "Shipping State",
        # Product Details (33 to 57: AG-BE)
        "Sl.No. *", "Product Description", "Is Service *", "HSN Code *", "Bar Code", "Quantity *", "Free Quantity", "Unit *", "Unit Price *",
        "Gross Amount", "Discount", "Pre Tax Value", "Taxable value *", "GST Rate (%) *", "Sgst Amt(Rs)", "Cgst Amt(Rs)",
        "Igst Amt(Rs)", "Cess Rate (%)", "Cess Amt Adval (Rs)", "Cess Non Adval Amt (Rs)", "State Cess Rate (%)",
        "State Cess Adval Amt (Rs)", "State Cess Non-Adval Amt (Rs)", "Other Charges", "Item Total *",
        # Spacer Cols (58 to 60: BF, BG, BH)
        "", "", "",
        # Value Details (61 to 70: BI-BR)
        "Total Taxable value *", "Sgst Amt", "Cgst Amt", "Igst Amt", "Cess Amt", "State Cess Amt", "Discount", "Other charges", "Round off", "Total Invoice value *",
        # Export Details (71 to 77: BS-BY)
        "Shipping Bill No", "Shipping Bill Dt", "Port", "Refund claim", "Foreign Currency", "Country Code", "Export Duty Amount",
        # E-way-bill Details (78 to 85: BZ-CG)
        "Trans ID", "Trans Name", "Trans Mode", "Distance", "Trans Doc No", "Trans Doc Date", "Vehicle No", "Vehicle Type"
    ]

    for c_idx, col_name in enumerate(col_names, 1):
        cell = ws.cell(row=4, column=c_idx, value=col_name)
        if c_idx <= 4:
            cell.fill = c_supply
        elif c_idx <= 7:
            cell.fill = c_doc
        elif c_idx <= 18:
            cell.fill = c_buyer
        elif c_idx <= 24:
            cell.fill = c_disp
        elif c_idx <= 32:
            cell.fill = c_ship
        elif c_idx <= 57:
            cell.fill = c_item
        elif 61 <= c_idx <= 70:
            cell.fill = c_val
        elif 71 <= c_idx <= 77:
            cell.fill = c_export
        elif 78 <= c_idx <= 85:
            cell.fill = c_eway
        cell.font = font_col
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    ws.row_dimensions[4].height = 42

    curr_row = 5
    for inv in data_list:
        is_interstate = inv["is_interstate"]
        pos_str = inv["pos_state_name"].title()

        item_rows_data = []
        tot_taxable = 0.0
        tot_cgst = 0.0
        tot_sgst = 0.0
        tot_igst = 0.0
        tot_discount = 0.0

        for it in inv["line_items"]:
            qty = it["qty"]
            price = it["unit_price"]
            gross_amt = it.get("gross_amt", round(qty * price, 2))
            disc = it.get("discount", 0.0)
            taxable = it.get("taxable_val", gross_amt)
            gst_rate = it.get("gst_rate", 5.0)

            tot_taxable += taxable
            tot_discount += disc

            if is_interstate:
                igst = round(taxable * (gst_rate / 100), 2)
                cgst = 0.0
                sgst = 0.0
                tot_igst += igst
            else:
                cgst = round(taxable * (gst_rate / 200), 2)
                sgst = round(taxable * (gst_rate / 200), 2)
                igst = 0.0
                tot_cgst += cgst
                tot_sgst += sgst

            item_tot = round(taxable + cgst + sgst + igst, 2)
            item_rows_data.append({
                "it": it, "qty": qty, "price": price, "gross_amt": gross_amt,
                "disc": disc, "taxable": taxable, "gst_rate": gst_rate,
                "cgst": cgst, "sgst": sgst, "igst": igst, "item_tot": item_tot
            })

        exact_total = round(tot_taxable + tot_cgst + tot_sgst + tot_igst, 2)
        final_inv_val = round(exact_total)
        exact_round_off = round(final_inv_val - exact_total, 2)

        if inv.get("shipping_is_different", False):
            ship_vals = [
                str(inv["ship_gstin"]), str(inv["ship_name"]), str(inv["ship_name"]),
                str(inv["ship_addr1"]), "", str(inv["ship_loc"]), str(inv["ship_pin"]), str(inv["ship_state_name"])
            ]
        else:
            ship_vals = ["", "", "", "", "", "", "", ""]

        for s_no, r_data in enumerate(item_rows_data, 1):
            it = r_data["it"]

            row_data = [
                # Supply Details (1 to 4: A-D)
                "B2B", "N", "", "N",
                # Document Details (5 to 7: E-G)
                "Tax Invoice", str(inv["invoice_no"]), str(inv["doc_date"]),
                # Buyer Details (8 to 18: H-R)
                str(inv["buyer_gstin"]), str(inv["buyer_name"]), str(inv["buyer_trade_name"]), pos_str,
                str(inv["buyer_addr1"]), "", str(inv["buyer_loc"]), str(inv["buyer_pin"]),
                str(inv["buyer_state_name"]), str(inv["buyer_phone"]), str(inv["buyer_email"]),
                # Dispatch Details (19 to 24: S-X)
                "", "", "", "", "", "",
                # Shipping Details (25 to 32: Y-AF)
                *ship_vals,
                # Product Details (33 to 57: AG-BE)
                str(s_no),                                    # AG (33): Sl.No. *
                str(it["desc"]),                              # AH (34): Product Description
                "N",                                          # AI (35): Is Service *
                str(it["hsn"]),                               # AJ (36): HSN Code *
                "",                                           # AK (37): Bar Code
                str(r_data["qty"]),                           # AL (38): Quantity *
                "0",                                          # AM (39): Free Quantity
                str(it.get("unit", "PAC")),                   # AN (40): Unit *
                fmt_dec(r_data["price"]),                     # AO (41): Unit Price *
                fmt_dec(r_data["gross_amt"]),                 # AP (42): Gross Amount
                fmt_dec(r_data["disc"]),                      # AQ (43): Discount (Item Level)
                "",                                           # AR (44): Pre Tax Value (BLANK AS REQUESTED)
                fmt_dec(r_data["taxable"]),                   # AS (45): Taxable value *
                str(int(r_data["gst_rate"])),                 # AT (46): GST Rate (%) *
                fmt_dec(r_data["sgst"]),                      # AU (47): Sgst Amt(Rs)
                fmt_dec(r_data["cgst"]),                      # AV (48): Cgst Amt(Rs)
                fmt_dec(r_data["igst"]),                      # AW (49): Igst Amt(Rs)
                "0", "0.00", "0.00", "0", "0.00", "0.00", "0.00", # AX to BD (50 to 56): Cess & Other
                fmt_dec(r_data["item_tot"]),                  # BE (57): Item Total *
                # Spacer Cols (58 to 60: BF, BG, BH)
                "", "", "",
                # Value Details (61 to 70: BI-BR)
                fmt_dec(tot_taxable),                         # BI (61): Total Taxable value *
                fmt_dec(tot_sgst),                            # BJ (62): Sgst Amt
                fmt_dec(tot_cgst),                            # BK (63): Cgst Amt
                fmt_dec(tot_igst),                            # BL (64): Igst Amt
                "0.00",                                       # BM (65): Cess Amt
                "0.00",                                       # BN (66): State Cess Amt
                fmt_dec(tot_discount),                        # BO (67): Discount (Invoice Total)
                "0.00",                                       # BP (68): Other charges
                fmt_dec(exact_round_off),                     # BQ (69): Round off (Strict Max ±0.99)
                fmt_dec(final_inv_val),                       # BR (70): Total Invoice value *
                # Export Details (71 to 77: BS-BY)
                "", "", "", "", "", "", "",
                # E-way-bill Details (78 to 85: BZ-CG)
                "", "", "", "", "", "", "", ""
            ]

            for c_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=curr_row, column=c_idx, value=str(val))
                cell.font = font_data
                cell.border = thin_border
                cell.number_format = '@'
                if c_idx in [1, 2, 4, 5, 7, 8, 11, 15, 16, 17, 25, 31, 32, 33, 35, 36, 38, 40, 46]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif c_idx in [41, 42, 43, 45, 47, 48, 49, 57, 61, 62, 63, 64, 67, 69, 70]:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

            curr_row += 1

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row < 3:
                continue
            v_str = str(cell.value or '')
            if len(v_str) > max_len:
                max_len = len(v_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

    # Hidden Columns matching utility view
    hidden_cols = [
        'B', 'C',                          # Reverse Charge, e-Comm GSTIN
        'N', 'O', 'P', 'Q', 'R',           # Buyer Location, Pin Code, State, Phone, Email
        'AI', 'AJ', 'AK', 'AL', 'AM', 'AN', 'AO', 'AP', 'AQ',  # Is Service to Discount
        'BF', 'BG', 'BH'                   # Spacer Buffer Columns
    ]
    for hc in hidden_cols:
        ws.column_dimensions[hc].hidden = True

    tabs = ["Welcome", "Profile", "Master Codes", "Sample Invoice", "Format A,B,C,D", "Schema", "Validation", "Calculations", "FAQs"]
    for t in tabs:
        d_ws = wb.create_sheet(title=t)
        d_ws.sheet_view.showGridLines = True
        d_ws["A1"] = f"{t} - Official e-Invoice System Utility"
        d_ws["A1"].font = Font(size=14, bold=True, color="1F497D")

    wb.active = ws

    out_io = io.BytesIO()
    wb.save(out_io)
    return out_io.getvalue()

# ----------------- CLEAN EXACT GOVERNMENT E-INVOICE PREVIEW -----------------

def render_exact_government_einvoice_preview(meta):
    is_interstate = meta["is_interstate"]

    tot_taxable = sum([it['taxable_val'] for it in meta['line_items']])
    tot_discount = sum([it.get('discount', 0.0) for it in meta['line_items']])

    if is_interstate:
        tot_igst = round(tot_taxable * 0.05, 2)
        tot_cgst = 0.0
        tot_sgst = 0.0
    else:
        tot_igst = 0.0
        tot_cgst = round(tot_taxable * 0.025, 2)
        tot_sgst = round(tot_taxable * 0.025, 2)

    exact_total = round(tot_taxable + tot_cgst + tot_sgst + tot_igst, 2)
    grand_total = round(exact_total)
    round_off = round(grand_total - exact_total, 2)
    ack_date_str = f"{meta['clean_date']} 17:0:00"

    table_rows = ""
    for idx, it in enumerate(meta['line_items'], 1):
        q = it['qty']
        p = it['unit_price']
        taxable = it['taxable_val']
        disc = it.get('discount', 0.0)
        gst_r = it.get('gst_rate', 5.0)

        tax_amt = round(taxable * (gst_r / 100), 2)
        tot_val = round(taxable + tax_amt, 2)

        table_rows += f"""
        <tr style="border-bottom: 1px solid #000; font-size: 11px;">
            <td style="border-right: 1px solid #000; padding: 5px 3px; text-align: center;">{idx}</td>
            <td style="border-right: 1px solid #000; padding: 5px 4px; text-align: left; font-weight: 500;">{it['desc']}</td>
            <td style="border-right: 1px solid #000; padding: 5px 3px; text-align: center;">{it['hsn']}</td>
            <td style="border-right: 1px solid #000; padding: 5px 3px; text-align: right;">{q}</td>
            <td style="border-right: 1px solid #000; padding: 5px 3px; text-align: center;">{it.get('unit', 'PAC')}</td>
            <td style="border-right: 1px solid #000; padding: 5px 3px; text-align: right;">{p:.2f}</td>
            <td style="border-right: 1px solid #000; padding: 5px 3px; text-align: right;">{disc:.2f}</td>
            <td style="border-right: 1px solid #000; padding: 5px 3px; text-align: right;">{taxable:.2f}</td>
            <td style="border-right: 1px solid #000; padding: 5px 3px; text-align: center; line-height: 1.2;">
                {gst_r:.2f}+0.00<br><span style="font-size: 9px; color: #555;">0.00+0</span>
            </td>
            <td style="border-right: 1px solid #000; padding: 5px 3px; text-align: right;">0</td>
            <td style="padding: 5px 4px; text-align: right; font-weight: bold;">{tot_val:.2f}</td>
        </tr>
        """

    html = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
body {{
    background-color: #f1f5f9;
    font-family: Arial, Helvetica, sans-serif;
    margin: 0;
    padding: 10px;
}}
.invoice-box {{
    background-color: #ffffff;
    color: #000000;
    padding: 16px;
    border: 2px solid #000;
    max-width: 930px;
    margin: 0 auto;
    box-shadow: 0 4px 10px rgba(0,0,0,0.15);
}}
.sec-title {{
    background-color: #d1d5db;
    color: #000;
    font-weight: bold;
    font-size: 11.5px;
    padding: 3px 6px;
    border: 1px solid #000;
}}
.sec-content {{
    border: 1px solid #000;
    border-top: none;
    padding: 6px;
    font-size: 11px;
    margin-bottom: 8px;
    line-height: 1.5;
}}
table.details {{
    width: 100%;
    border-collapse: collapse;
}}
table.details td {{
    padding: 2px 4px;
    font-size: 11px;
}}
table.items {{
    width: 100%;
    border-collapse: collapse;
    border: 1px solid #000;
    border-top: none;
}}
table.items th {{
    background-color: #f3f4f6;
    border-right: 1px solid #000;
    border-bottom: 1px solid #000;
    padding: 5px 3px;
    font-size: 10.5px;
}}
</style>
</head>
<body>
<div class="invoice-box">
    <!-- TOP HEADER -->
    <table style="width: 100%; border-collapse: collapse; margin-bottom: 8px;">
        <tr>
            <td style="vertical-align: top; width: 75%;">
                <div style="font-size: 19px; font-weight: bold; letter-spacing: 0.5px;">{meta['seller_gstin']}</div>
                <div style="font-size: 17px; font-weight: bold; margin-top: 4px; text-transform: uppercase;">{meta['seller_name']}</div>
            </td>
            <td style="vertical-align: top; width: 25%; text-align: right;">
                <div style="display: inline-block; width: 95px; height: 95px; border: 1.5px solid #000; text-align: center; background: #fff; padding: 4px;">
                    <svg viewBox="0 0 100 100" style="width: 100%; height: 100%;">
                        <rect width="100" height="100" fill="#fff" />
                        <path d="M10 10 h30 v30 h-30 z M15 15 v20 h20 v-20 z M20 20 h10 v10 h-10 z" fill="#000" />
                        <path d="M60 10 h30 v30 h-30 z M65 15 v20 h20 v-20 z M70 20 h10 v10 h-10 z" fill="#000" />
                        <path d="M10 60 h30 v30 h-30 z M15 65 v20 h20 v-20 z M20 70 h10 v10 h-10 z" fill="#000" />
                        <rect x="50" y="50" width="10" height="10" fill="#000" />
                        <rect x="65" y="65" width="15" height="15" fill="#000" />
                        <rect x="75" y="50" width="10" height="10" fill="#000" />
                        <rect x="50" y="75" width="10" height="10" fill="#000" />
                    </svg>
                </div>
            </td>
        </tr>
    </table>

    <!-- SECTION 1: e-Invoice Details -->
    <div class="sec-title">1. e-Invoice Details</div>
    <div class="sec-content">
        <table style="width: 100%;">
            <tr>
                <td style="width: 50%;"><b>IRN :</b> <span style="font-size: 9.5px; word-break: break-all;">3521a723ac0d702f87a9ee33b47e4f25eb9b8ac3e0d2160434014ee0c0102180</span></td>
                <td style="width: 25%;"><b>Ack. No :</b> 172621231914553</td>
                <td style="width: 25%; text-align: right;"><b>Ack. Date :</b> {ack_date_str}</td>
            </tr>
        </table>
    </div>

    <!-- SECTION 2: Transaction Details -->
    <div class="sec-title">2. Transaction Details</div>
    <div class="sec-content" style="line-height: 1.6;">
        <table style="width: 100%;">
            <tr>
                <td style="width: 32%;"><b>Supply Type Code :</b> B2B</td>
                <td style="width: 35%;"><b>Document No :</b> {meta['invoice_no']}</td>
                <td style="width: 33%;" rowspan="2"><b>IGST applicable despite Supplier and Recipient located in same State :</b> No</td>
            </tr>
            <tr>
                <td><b>Place of Supply :</b> {meta['pos_state_name'].upper()}</td>
                <td></td>
            </tr>
            <tr>
                <td><b>Document Type :</b> Tax Invoice</td>
                <td><b>Document Date :</b> {meta['clean_date']}</td>
                <td></td>
            </tr>
        </table>
    </div>

    <!-- SECTION 3: Party Details -->
    <div class="sec-title">3. Party Details</div>
    <div style="border: 1px solid #000; border-top: none; margin-bottom: 8px;">
        <table style="width: 100%; border-collapse: collapse; font-size: 11px;">
            <tr>
                <td style="width: 50%; border-right: 1px solid #000; padding: 6px; vertical-align: top;">
                    <div style="font-weight: bold; font-size: 13px; text-decoration: underline; margin-bottom: 3px;">Supplier</div>
                    <div><b>GSTIN :</b> {meta['seller_gstin']}</div>
                    <div style="font-weight: bold; margin-top: 2px;">{meta['seller_name']}</div>
                    <div>{meta['seller_addr']}</div>
                    <div>{meta['seller_loc'].upper()}</div>
                    <div>{meta['seller_pin']} &nbsp; {meta['seller_state_name'].upper()}</div>
                    <div style="margin-top: 3px;">+91-7070701513 &nbsp; info@romsons.in</div>
                </td>
                <td style="width: 50%; padding: 6px; vertical-align: top;">
                    <div style="font-weight: bold; font-size: 13px; text-decoration: underline; margin-bottom: 3px;">Recipient</div>
                    <div><b>GSTIN :</b> {meta['buyer_gstin']}</div>
                    <div style="font-weight: bold; margin-top: 2px;">{meta['buyer_name']}</div>
                    <div>{meta['buyer_addr1']}</div>
                    <div>{meta['buyer_loc']} &nbsp; Place of Supply: {meta['pos_state_name'].upper()}</div>
                    <div>{meta['buyer_pin']} &nbsp; {meta['buyer_state_name'].upper()}</div>
                    <div style="margin-top: 3px;">{meta['buyer_phone']} &nbsp; {meta['buyer_email']}</div>
                </td>
            </tr>
        </table>
    </div>

    <!-- SECTION 4: Details of Goods / Services -->
    <div class="sec-title">4. Details of Goods / Services</div>
    <table class="items">
        <thead>
            <tr>
                <th style="width: 4%;">SlNo</th>
                <th style="width: 28%; text-align: left;">Item Description</th>
                <th style="width: 8%;">HSN Code</th>
                <th style="width: 6%;">Quantity</th>
                <th style="width: 5%;">Unit</th>
                <th style="width: 8%;">Unit Price(Rs)</th>
                <th style="width: 7%;">Discount(Rs)</th>
                <th style="width: 9%;">Taxable Amount(Rs)</th>
                <th style="width: 14%; line-height: 1.2;">
                    Tax Rate<br><span style="font-size: 9px; font-weight: normal;">(GST+Cess | State Cess+Cess Non.Advol)</span>
                </th>
                <th style="width: 6%;">Other charges(Rs)</th>
                <th style="width: 9%; border-right: none;">Total</th>
            </tr>
        </thead>
        <tbody>
            {table_rows}
        </tbody>
    </table>

    <!-- TOTALS & SUMMARY -->
    <table style="width: 100%; border-collapse: collapse; border: 1px solid #000; border-top: none; font-size: 11px; margin-top: -1px;">
        <tr>
            <td style="width: 60%; padding: 6px; vertical-align: top; border-right: 1px solid #000; font-size: 10px; color: #444;">
                * This is an exact live verification preview matching the official e-Invoice standard format. Official IRN & QR Code will be issued upon uploading the generated JSON to the IRP portal.
            </td>
            <td style="width: 40%; padding: 6px; vertical-align: top;">
                <table style="width: 100%; line-height: 1.6;">
                    <tr>
                        <td>Total Taxable Value :</td>
                        <td style="text-align: right; font-weight: bold;">₹{tot_taxable:.2f}</td>
                    </tr>
                    <tr>
                        <td>Total Discount :</td>
                        <td style="text-align: right; color: #dc2626;">₹{tot_discount:.2f}</td>
                    </tr>
                    <tr>
                        <td>Total CGST :</td>
                        <td style="text-align: right;">₹{tot_cgst:.2f}</td>
                    </tr>
                    <tr>
                        <td>Total SGST :</td>
                        <td style="text-align: right;">₹{tot_sgst:.2f}</td>
                    </tr>
                    <tr>
                        <td>Total IGST :</td>
                        <td style="text-align: right;">₹{tot_igst:.2f}</td>
                    </tr>
                    <tr>
                        <td>Round Off :</td>
                        <td style="text-align: right;">₹{round_off:.2f}</td>
                    </tr>
                    <tr style="border-top: 1.5px solid #000; font-size: 13px; font-weight: bold;">
                        <td>Total Invoice Value :</td>
                        <td style="text-align: right;">₹{grand_total:.2f}</td>
                    </tr>
                </table>
            </td>
        </tr>
    </table>
</div>
</body>
</html>
"""
    return html

# ----------------- MAIN STREAMLIT WORKFLOW -----------------

if uploaded_invoices:
    processed_docs = []
    for file in uploaded_invoices:
        try:
            pdf_bytes = file.read()
            reconciled_pdf_buf, meta = process_and_reconcile_pdf(pdf_bytes)
            processed_docs.append({
                "pdf_buf": reconciled_pdf_buf,
                "meta": meta
            })
        except Exception as e:
            st.error(f"Error processing {file.name}: {str(e)}")

    if processed_docs:
        st.success(f"✅ Successfully processed {len(processed_docs)} Invoices!")
        
        parsed_invoices = [doc["meta"] for doc in processed_docs]
        excel_raw_bytes = generate_official_nic_v101_excel_bytes(parsed_invoices)

        # SINGLE INVOICE ACTIONS
        if len(processed_docs) == 1:
            doc = processed_docs[0]
            meta = doc["meta"]
            pdf_buf = doc["pdf_buf"]
            json_payload = build_einvoice_json_v101(meta)
            inv_no = meta["invoice_no"]
            pdf_name = meta["pdf_filename"]
            is_blinkit = meta["is_blinkit"]

            channel_badge = "🟢 Blinkit Conversion Active (50x/25x Box Fix)" if is_blinkit else "🔵 Standard Buyer (Original Data Preserved)"
            ship_status = "⚠️ Different (Shipping Details Filled)" if meta["shipping_is_different"] else "✅ Same (Shipping Blank in Excel)"

            st.markdown(f"""
                <div style="background-color: #1e293b; padding: 14px 18px; border-radius: 8px; margin-bottom: 15px; border-left: 5px solid {'#10b981' if is_blinkit else '#3b82f6'};">
                    <b>Buyer Name:</b> <code>{meta['buyer_name']}</code> &nbsp;|&nbsp; <b>Rule:</b> <code>{channel_badge}</code><br>
                    <b>City:</b> <code>{meta['buyer_loc']}</code> ({meta['buyer_pin']}) &nbsp;|&nbsp; <b>Shipping Status:</b> <code>{ship_status}</code><br>
                    <b>Seller:</b> {meta['seller_name']} ({meta['seller_gstin']})
                </div>
            """, unsafe_allow_html=True)

            # DOWNLOAD BUTTONS
            c1, c2, c3 = st.columns(3)
            with c1:
                st.download_button(
                    label=f"📥 Download Processed PDF",
                    data=pdf_buf.getvalue(),
                    file_name=pdf_name,
                    mime="application/pdf"
                )
            with c2:
                st.download_button(
                    label=f"📊 Download eInvoice.xlsx (Exact Template)",
                    data=excel_raw_bytes,
                    file_name=f"{inv_no}_eInvoice.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
            with c3:
                st.download_button(
                    label=f"🧾 Download e-Invoice JSON (v1.01)",
                    data=json.dumps(json_payload, indent=4),
                    file_name=f"{inv_no}_eInvoice_v1.01.json",
                    mime="application/json"
                )

            # LIVE GOVERNMENT FORMAT PREVIEW
            st.markdown("<br>", unsafe_allow_html=True)
            st.markdown("### 📄 Official Government e-Invoice Live Preview")
            
            calc_height = 560 + (len(meta['line_items']) * 36)
            preview_html = render_exact_government_einvoice_preview(meta)
            components.html(preview_html, height=calc_height, scrolling=True)

        # BULK INVOICE ACTIONS
        else:
            zip_buffer = io.BytesIO()
            with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
                for idx, doc in enumerate(processed_docs, 1):
                    pdf_name = doc["meta"]["pdf_filename"]
                    zip_file.writestr(pdf_name, doc["pdf_buf"].getvalue())
            zip_buffer.seek(0)

            bulk_json = [build_einvoice_json_v101(d["meta"]) for d in processed_docs]
            date_str = datetime.now().strftime('%d-%m-%Y')

            c1, c2, c3 = st.columns(3)
            with c1:
                st.download_button(
                    label=f"📦 Download All Processed PDFs ({len(processed_docs)} Files - ZIP)",
                    data=zip_buffer.getvalue(),
                    file_name=f"Invoices_Processed_{date_str}.zip",
                    mime="application/zip"
                )
            with c2:
                st.download_button(
                    label=f"📊 Download Bulk eInvoice.xlsx ({len(processed_docs)} Invoices)",
                    data=excel_raw_bytes,
                    file_name=f"Bulk_eInvoice_{date_str}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
            with c3:
                st.download_button(
                    label=f"🧾 Download Bulk e-Invoice JSON ({len(bulk_json)} Invoices)",
                    data=json.dumps(bulk_json, indent=4),
                    file_name=f"Bulk_eInvoice_NIC_v1.01_{date_str}.json",
                    mime="application/json"
                )

            st.markdown("<br>", unsafe_allow_html=True)
            st.markdown("### 📄 Invoices Official e-Invoice Preview")
            preview_tabs = st.tabs([f"INV: {d['meta']['invoice_no']}" for d in processed_docs])
            for i, tab in enumerate(preview_tabs):
                with tab:
                    curr_meta = processed_docs[i]['meta']
                    calc_height = 560 + (len(curr_meta['line_items']) * 36)
                    p_html = render_exact_government_einvoice_preview(curr_meta)
                    components.html(p_html, height=calc_height, scrolling=True)
